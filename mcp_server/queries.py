"""Read-only queries behind the Career OS MCP tools (Fase 1).

Pure functions over the repo's JSON files, kept apart from the MCP wiring so they
are testable without a client. Nothing here writes to disk.
"""

from __future__ import annotations

import csv
import datetime
import re
from pathlib import Path
from typing import Any, Iterable, Optional

from engine.automation.dedupe import find_duplicate, normalize_tokens
from job_store import ROOT, list_jobs, read_json
from resume_ids import build_index

MASTER_PATH = ROOT / "master_resume.json"
APPLICATIONS_PATH = ROOT / "applications.json"

LANGUAGES = ("pt", "en")
EVIDENCE_KINDS = ("bullet", "responsibility", "role", "language", "education", "certification", "volunteer")
SEM_RESPOSTA_DIAS_PADRAO = 14


# --- master resume ------------------------------------------------------------

def _pick_language(value: Any, idioma: str) -> Any:
    """Keep only `idioma` wherever a dict carries both pt and en."""
    if isinstance(value, dict):
        if "pt" in value and "en" in value:
            value = {k: v for k, v in value.items() if k not in LANGUAGES or k == idioma}
        return {k: _pick_language(v, idioma) for k, v in value.items()}
    if isinstance(value, list):
        return [_pick_language(item, idioma) for item in value]
    return value


def master_evidence(master: dict[str, Any], idioma: str = "pt") -> list[dict[str, Any]]:
    """Flat list of citable items — what an evidence map points at.

    Responsibilities that restate a bullet (`same_as`) are left out: cite the bullet."""
    companies = {exp["id"]: exp.get("company") for exp in master.get("work_experience", [])}
    items = []
    for item_id, entry in build_index(master).items():
        kind = entry["kind"]
        if kind not in EVIDENCE_KINDS or entry.get("same_as"):
            continue
        if kind == "role":
            text = (entry.get("title") or {}).get(idioma)
        elif kind == "language":
            text = f"{entry.get('language')}: {entry.get(idioma)}"
        elif kind == "education":
            text = f"{entry.get(idioma)} — {entry.get('institution')}"
        elif kind == "certification":
            text = f"{entry.get('name')} ({entry.get('issuer')}) — {entry.get(idioma)}"
        elif kind == "volunteer":
            text = f"{entry.get('organization')}: {entry.get(idioma)}"
        else:
            text = entry.get(idioma)
        items.append({
            "id": item_id,
            "tipo": kind,
            "empresa": companies.get(entry.get("experience_id")),
            "texto": text,
        })
    return items


def read_master(formato: str = "evidencias", idioma: str = "pt",
                secoes: Optional[list[str]] = None) -> dict[str, Any]:
    master = read_json(MASTER_PATH)
    if idioma not in (*LANGUAGES, "ambos"):
        raise ValueError("idioma deve ser pt, en ou ambos")
    if formato == "evidencias":
        languages = LANGUAGES if idioma == "ambos" else (idioma,)
        return {lang: master_evidence(master, lang) for lang in languages}
    if formato != "completo":
        raise ValueError("formato deve ser evidencias ou completo")
    if secoes:
        unknown = sorted(set(secoes) - set(master))
        if unknown:
            raise ValueError(f"seções inexistentes: {unknown}. Disponíveis: {sorted(master)}")
        master = {key: master[key] for key in secoes}
    return master if idioma == "ambos" else _pick_language(master, idioma)


# --- new jobs -----------------------------------------------------------------

_COLUMN_ALIASES = {
    "empresa": ("empresa", "company", "company_name", "companhia"),
    "cargo": ("cargo", "vaga", "role", "title", "titulo", "título", "role_title", "posicao", "posição"),
    "url": ("url", "link", "link da vaga", "job_url"),
    "fit": ("fit", "fit_score", "score", "match", "aderencia", "aderência"),
    "aplicada": ("aplicada", "applied", "candidatura enviada"),
    "status_planilha": ("status no radar", "status", "situação", "situacao"),
}
_CLOSED_WORDS = ("encerrad", "fechad", "expirad", "descartad", "closed")


def _canonical(row: dict[str, Any]) -> dict[str, Any]:
    lowered = {str(k).strip().casefold(): v for k, v in row.items() if k is not None}
    out: dict[str, Any] = {}
    for field, aliases in _COLUMN_ALIASES.items():
        out[field] = next((lowered[a] for a in aliases if lowered.get(a) not in (None, "")), None)
    out["fit"] = _parse_fit(out["fit"])
    out["aplicada"] = str(out["aplicada"]).strip().casefold() in {"true", "sim", "yes", "x", "1", "verdadeiro"}
    for field in ("empresa", "cargo", "url", "status_planilha"):
        out[field] = str(out[field]).strip() if out[field] is not None else ""
    out["dados"] = {str(k): v for k, v in row.items() if k is not None}
    return out


def _parse_fit(value: Any) -> Optional[float]:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        number = float(value)
    else:
        match = re.search(r"\d+(?:[.,]\d+)?", str(value))
        if not match:
            return None
        number = float(match.group().replace(",", "."))
    # Normalize to 0-100: 0.93 and 9.3 (0-10 scale, as in the Radar sheet) both mean 93.
    if 0 < number <= 1:
        return number * 100
    if number <= 10:
        return number * 10
    return number


def read_sheet(path: Path) -> list[dict[str, Any]]:
    """First sheet of an .xlsx, or a .csv; header on the first row."""
    if path.suffix.lower() == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as file:
            return list(csv.DictReader(file))
    from openpyxl import load_workbook

    sheet = load_workbook(path, read_only=True, data_only=True).worksheets[0]
    rows = sheet.iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else None for h in next(rows, [])]
    return [dict(zip(header, values)) for values in rows if any(v not in (None, "") for v in values)]


def classify_jobs(rows: Iterable[dict[str, Any]], jobs: list[dict[str, Any]],
                  applications: list[dict[str, Any]]) -> dict[str, list[dict[str, Any]]]:
    """Split incoming rows into novas / ja_adaptadas / descartadas / ja_aplicadas."""
    discarded_ids = {job["id"] for job in jobs if job.get("triage", {}).get("decision") == "discard"
                     or job.get("automation", {}).get("workflow_status") in {"DISCARDED", "EXPIRED"}}
    result: dict[str, list[dict[str, Any]]] = {"novas": [], "ja_adaptadas": [], "descartadas": [],
                                               "ja_aplicadas": [], "incompletas": []}
    open_by_company: dict[frozenset, list[dict[str, Any]]] = {}
    for app in applications:
        if app.get("status") not in {"Rejeitado", "Desistente"}:
            open_by_company.setdefault(frozenset(normalize_tokens(app.get("company", ""))), []).append(app)
    seen_in_batch: set[tuple[str, str]] = set()
    for raw in rows:
        row = _canonical(raw)
        if not row["empresa"] and not row["cargo"]:
            continue  # blank spreadsheet row
        if not row["empresa"] or not row["cargo"]:
            result["incompletas"].append(row)
            continue
        key = (row["empresa"].casefold(), row["cargo"].casefold())
        if key in seen_in_batch:
            continue
        seen_in_batch.add(key)

        if row["aplicada"]:
            result["ja_aplicadas"].append({**row, "match": "planilha: Aplicada"})
            continue
        if any(word in row["status_planilha"].casefold() for word in _CLOSED_WORDS):
            result["descartadas"].append({**row, "match": f"planilha: {row['status_planilha']}"})
            continue
        applied = find_duplicate(url=row["url"], company=row["empresa"], role=row["cargo"],
                                 existing_jobs=[], existing_applications=applications)
        if applied:
            result["ja_aplicadas"].append({**row, "match": applied})
            continue
        known = find_duplicate(url=row["url"], company=row["empresa"], role=row["cargo"],
                               existing_jobs=jobs, existing_applications=[])
        if known in discarded_ids:
            result["descartadas"].append({**row, "match": known})
        elif known:
            result["ja_adaptadas"].append({**row, "match": known})
        else:
            same_company = open_by_company.get(frozenset(normalize_tokens(row["empresa"])), [])
            if same_company:
                row["atencao"] = "Candidatura em aberto na mesma empresa: " + "; ".join(
                    f"#{a.get('id')} {a.get('role')} ({a.get('status')})" for a in same_company)
            result["novas"].append(row)

    result["novas"].sort(key=lambda r: (r["fit"] is None, -(r["fit"] or 0)))
    result["ja_adaptadas"].sort(key=lambda r: (r["fit"] is None, -(r["fit"] or 0)))
    return result


def new_jobs(vagas: Optional[list[dict[str, Any]]] = None, planilha: Optional[str] = None,
             limite: int = 30) -> dict[str, Any]:
    if vagas is None:
        if not planilha:
            raise ValueError(
                "Informe `vagas` (linhas já lidas da planilha, ex.: via conector do Google Drive) "
                "ou `planilha` / variável CAREER_OS_PLANILHA com o caminho de um .xlsx ou .csv local."
            )
        path = Path(planilha).expanduser()
        if not path.is_absolute():
            path = ROOT / path
        if not path.exists():
            raise ValueError(f"planilha não encontrada: {path}")
        vagas = read_sheet(path)

    classified = classify_jobs(vagas, list_jobs(), read_json(APPLICATIONS_PATH))
    slim = lambda rows: [{k: r[k] for k in ("empresa", "cargo", "url", "fit", "match") if k in r} for r in rows]
    for row in classified["novas"]:
        row.pop("aplicada", None)
    return {
        "resumo": {k: len(v) for k, v in classified.items()},
        "novas": classified["novas"][:limite],
        "ja_adaptadas": slim(classified["ja_adaptadas"]),
        "descartadas": slim(classified["descartadas"]),
        "ja_aplicadas": slim(classified["ja_aplicadas"]),
        "incompletas": [r["dados"] for r in classified["incompletas"]],
    }


# --- applications -------------------------------------------------------------

def _last_movement(app: dict[str, Any]) -> Optional[str]:
    dated = [h["date"] for h in app.get("status_history", []) if h.get("date")]
    return max(dated) if dated else (app.get("date_applied") or None)


def list_applications(status: Optional[list[str]] = None, sem_resposta: bool = False,
                      sem_resposta_dias: int = SEM_RESPOSTA_DIAS_PADRAO, busca: str = "",
                      desde: Optional[str] = None, hoje: Optional[datetime.date] = None) -> dict[str, Any]:
    today = hoje or datetime.date.today()
    applications = read_json(APPLICATIONS_PATH)
    rows = []
    for app in applications:
        last = _last_movement(app)
        idle = (today - datetime.date.fromisoformat(last)).days if last else None
        rows.append({
            "id": app.get("id"),
            "empresa": app.get("company"),
            "cargo": app.get("role"),
            "status": app.get("status"),
            "etapa_atual": app.get("current_stage"),
            "data_candidatura": app.get("date_applied"),
            "ultima_movimentacao": last,
            "dias_parado": idle,
            "sem_resposta": app.get("status") == "Candidatado" and idle is not None and idle >= sem_resposta_dias,
            "fit": app.get("fit_score"),
            "url": app.get("url"),
            "slug_vaga": app.get("source_job_id"),
            "historico": app.get("status_history", []),
        })

    funnel: dict[str, int] = {}
    for row in rows:
        funnel[row["status"]] = funnel.get(row["status"], 0) + 1
    funnel["sem_resposta (derivado)"] = sum(r["sem_resposta"] for r in rows)

    if status:
        wanted = {s.casefold() for s in status}
        rows = [r for r in rows if (r["status"] or "").casefold() in wanted]
    if sem_resposta:
        rows = [r for r in rows if r["sem_resposta"]]
    if desde:
        rows = [r for r in rows if (r["data_candidatura"] or "") >= desde]
    if busca:
        needle = busca.casefold()
        rows = [r for r in rows if needle in f"{r['empresa']} {r['cargo']}".casefold()]
    rows.sort(key=lambda r: r["ultima_movimentacao"] or "", reverse=True)
    return {"funil": funnel, "total_filtrado": len(rows), "candidaturas": rows}


# --- saved jobs ---------------------------------------------------------------

def read_job(slug: str, idioma: str = "ambos") -> dict[str, Any]:
    from job_store import job_path, load_job

    if not job_path(slug).exists():
        known = sorted(job["id"] for job in list_jobs())
        raise ValueError(f"vaga não encontrada: {slug}. Salvas: {known}")
    job = load_job(slug)
    return job if idioma == "ambos" else _pick_language(job, idioma)

